from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


PARSER_VERSION = "1.0.0"


def parse_document(
    *,
    document_id: str,
    original_filename: str,
    file_format: str,
    content: bytes,
) -> dict[str, Any]:
    extension = Path(original_filename).suffix.lower()
    normalized_format = file_format.upper()
    if normalized_format == "PDF" or extension == ".pdf":
        return parse_pdf(document_id=document_id, content=content)
    if normalized_format == "EXCEL" or extension in {".xlsx", ".xls"}:
        if extension == ".xls":
            return parse_xls(document_id=document_id, content=content)
        return parse_xlsx(document_id=document_id, content=content)
    raise ValueError(f"Unsupported preview format for this phase: {original_filename}")


def serialize_preview(preview: dict[str, Any], max_size_bytes: int) -> bytes:
    raw = json.dumps(preview, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    if len(raw) > max_size_bytes:
        raise ValueError("Generated preview exceeds configured JSON size limit")
    return raw


def parse_pdf(*, document_id: str, content: bytes) -> dict[str, Any]:
    import fitz

    try:
        pdf = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("Invalid or corrupt PDF file") from exc
    if pdf.needs_pass:
        raise ValueError("Password-protected PDF is not supported in preview phase")

    pages: list[dict[str, Any]] = []
    total_text_blocks = 0
    total_tables = 0
    for index, page in enumerate(pdf, start=1):
        page_dict = page.get_text("dict")
        text_blocks: list[dict[str, Any]] = []
        lines: list[dict[str, Any]] = []
        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            block_lines = block.get("lines", [])
            block_text = " ".join(
                span.get("text", "")
                for line in block_lines
                for span in line.get("spans", [])
                if span.get("text")
            ).strip()
            if block_text:
                text_blocks.append(
                    {
                        "block_id": str(uuid4()),
                        "text": block_text,
                        "bbox": _bbox(block.get("bbox")),
                        "confidence": None,
                        "source": "pdf_text_layer",
                    }
                )
            for line in block_lines:
                tokens = []
                line_text_parts = []
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    if not text:
                        continue
                    line_text_parts.append(text)
                    tokens.append(
                        {
                            "token_id": str(uuid4()),
                            "text": text,
                            "bbox": _bbox(span.get("bbox")),
                            "confidence": None,
                        }
                    )
                line_text = " ".join(part.strip() for part in line_text_parts if part.strip()).strip()
                if line_text:
                    lines.append(
                        {
                            "line_id": str(uuid4()),
                            "text": line_text,
                            "bbox": _bbox(line.get("bbox")),
                            "tokens": tokens,
                        }
                    )
        tables = _detect_pdf_table_candidates(lines)
        total_text_blocks += len(text_blocks)
        total_tables += len(tables)
        pages.append(
            {
                "page_number": index,
                "width": float(page.rect.width),
                "height": float(page.rect.height),
                "rotation": int(page.rotation),
                "text_blocks": text_blocks,
                "lines": lines,
                "tables": tables,
            }
        )

    requires_ocr = total_text_blocks == 0
    return {
        "document_id": document_id,
        "file_format": "PDF",
        "parser_version": PARSER_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "requires_ocr": requires_ocr,
        "ocr_reason": "no_text_layer_detected" if requires_ocr else None,
        "pages": pages,
        "summary": {
            "page_count": len(pages),
            "sheet_count": 0,
            "text_block_count": total_text_blocks,
            "table_count": total_tables,
        },
    }


def parse_xlsx(*, document_id: str, content: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Invalid or corrupt Excel file") from exc

    sheets = []
    total_tables = 0
    for index, sheet in enumerate(workbook.worksheets):
        sheet_preview = _parse_openpyxl_sheet(sheet, index)
        total_tables += len(sheet_preview["detected_tables"])
        sheets.append(sheet_preview)
    return {
        "document_id": document_id,
        "file_format": "XLSX",
        "parser_version": PARSER_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "requires_ocr": False,
        "sheets": sheets,
        "summary": {
            "page_count": 0,
            "sheet_count": len(sheets),
            "text_block_count": 0,
            "table_count": total_tables,
        },
    }


def parse_xls(*, document_id: str, content: bytes) -> dict[str, Any]:
    try:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise ValueError("Invalid or corrupt XLS file") from exc

    sheets = []
    total_tables = 0
    for index in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(index)
        cells = []
        non_empty: list[tuple[int, int]] = []
        for row_idx in range(sheet.nrows):
            for col_idx in range(sheet.ncols):
                value = sheet.cell_value(row_idx, col_idx)
                if value in ("", None):
                    continue
                row = row_idx + 1
                col = col_idx + 1
                non_empty.append((row, col))
                cells.append(
                    {
                        "cell_id": str(uuid4()),
                        "address": f"{get_column_letter(col)}{row}",
                        "row": row,
                        "column": col,
                        "column_letter": get_column_letter(col),
                        "value": value,
                        "raw_value": value,
                        "data_type": _excel_value_type(value),
                        "is_merged": False,
                        "merged_range": None,
                    }
                )
        detected_headers, detected_tables = _detect_excel_regions(non_empty, cells)
        total_tables += len(detected_tables)
        sheets.append(
            {
                "sheet_id": str(uuid4()),
                "name": sheet.name,
                "index": index,
                "max_row": sheet.nrows,
                "max_column": sheet.ncols,
                "cells": cells,
                "merged_cells": [],
                "detected_headers": detected_headers,
                "detected_tables": detected_tables,
            }
        )
    return {
        "document_id": document_id,
        "file_format": "XLS",
        "parser_version": PARSER_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "requires_ocr": False,
        "sheets": sheets,
        "summary": {
            "page_count": 0,
            "sheet_count": len(sheets),
            "text_block_count": 0,
            "table_count": total_tables,
        },
    }


def _parse_openpyxl_sheet(sheet, index: int) -> dict[str, Any]:
    merged_ranges = [_merged_range_payload(str(item)) for item in sheet.merged_cells.ranges]
    merged_by_cell: dict[str, str] = {}
    for item in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(item))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_by_cell[f"{get_column_letter(col)}{row}"] = str(item)

    cells = []
    non_empty: list[tuple[int, int]] = []
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            address = cell.coordinate
            if value in ("", None) and address not in merged_by_cell:
                continue
            if value not in ("", None):
                non_empty.append((cell.row, cell.column))
            cells.append(
                {
                    "cell_id": str(uuid4()),
                    "address": address,
                    "row": cell.row,
                    "column": cell.column,
                    "column_letter": get_column_letter(cell.column),
                    "value": value,
                    "raw_value": value,
                    "data_type": _excel_value_type(value),
                    "is_merged": address in merged_by_cell,
                    "merged_range": merged_by_cell.get(address),
                }
            )

    detected_headers, detected_tables = _detect_excel_regions(non_empty, cells)
    return {
        "sheet_id": str(uuid4()),
        "name": sheet.title,
        "index": index,
        "max_row": sheet.max_row or 0,
        "max_column": sheet.max_column or 0,
        "cells": cells,
        "merged_cells": merged_ranges,
        "detected_headers": detected_headers,
        "detected_tables": detected_tables,
    }


def _bbox(value: Any) -> dict[str, float]:
    x0, y0, x1, y1 = value or (0.0, 0.0, 0.0, 0.0)
    return {"x0": float(x0), "y0": float(y0), "x1": float(x1), "y1": float(y1)}


def _merged_range_payload(value: str) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = range_boundaries(value)
    return {
        "range": value,
        "start_row": min_row,
        "start_column": min_col,
        "end_row": max_row,
        "end_column": max_col,
    }


def _excel_value_type(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if hasattr(value, "isoformat"):
        return "date"
    return "text"


def _detect_excel_regions(
    non_empty: list[tuple[int, int]],
    cells: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not non_empty:
        return [], []
    by_row: dict[int, list[dict[str, Any]]] = {}
    for cell in cells:
        if cell["value"] in ("", None):
            continue
        by_row.setdefault(cell["row"], []).append(cell)

    header_row = None
    for row_number in sorted(by_row):
        row_cells = by_row[row_number]
        text_values = [str(cell["value"]) for cell in row_cells if cell["data_type"] == "text"]
        if len(row_cells) >= 2 and len(text_values) >= max(1, len(row_cells) // 2):
            header_row = row_number
            break

    detected_headers = []
    if header_row is not None:
        detected_headers.append(
            {
                "row": header_row,
                "values": [str(cell["value"]) for cell in sorted(by_row[header_row], key=lambda item: item["column"])],
                "confidence": 0.85,
            }
        )

    min_row = min(row for row, _ in non_empty)
    max_row = max(row for row, _ in non_empty)
    min_col = min(col for _, col in non_empty)
    max_col = max(col for _, col in non_empty)
    if header_row is None:
        header_row = min_row
    table_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    return detected_headers, [
        {
            "table_id": str(uuid4()),
            "range": table_range,
            "header_row": header_row,
            "start_row": min(header_row + 1, max_row),
            "end_row": max_row,
            "confidence": 0.8 if max_row > min_row and max_col > min_col else 0.45,
        }
    ]


def _detect_pdf_table_candidates(lines: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates = [
        line for line in lines if len(line.get("tokens", [])) >= 3 or len(line.get("text", "").split()) >= 3
    ]
    if len(candidates) < 3:
        return []
    x0 = min(line["bbox"]["x0"] for line in candidates)
    y0 = min(line["bbox"]["y0"] for line in candidates)
    x1 = max(line["bbox"]["x1"] for line in candidates)
    y1 = max(line["bbox"]["y1"] for line in candidates)
    return [
        {
            "table_id": str(uuid4()),
            "bbox": {"x0": x0, "y0": y0, "x1": x1, "y1": y1},
            "rows": [],
            "columns": [],
            "confidence": 0.55,
        }
    ]
